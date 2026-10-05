"""Export the Design weekly report to a professional .docx."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
OUT_DIR = Path(__file__).resolve().parent

NAVY = "1B365D"
NAVY_RGB = RGBColor(0x1B, 0x36, 0x5D)
GOLD = "B08D57"
INK = RGBColor(0x2B, 0x2B, 0x2B)
MUTED = RGBColor(0x5C, 0x64, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = "F4F6F8"
ROW_OK = "E8F2EC"
ROW_INFO = "E8EEF6"
HEADER_FILL = NAVY
LINE = "D0D5DD"
WEEK_START = date(2026, 9, 28)
WEEK_END = date(2026, 10, 4)
AUDIT = date(2026, 10, 4)


def qn_w(tag: str) -> str:
    return qn(tag)


def set_run_font(run, name="Calibri", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn_w("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn_w("w:ascii"), name)
    rFonts.set(qn_w("w:hAnsi"), name)
    rFonts.set(qn_w("w:eastAsia"), name)
    rFonts.set(qn_w("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn_w("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn_w("w:val"), "clear")
    shd.set(qn_w("w:color"), "auto")
    shd.set(qn_w("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn_w("w:w"), str(value))
        node.set(qn_w("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color=LINE, sz="4") -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)
    for child in list(tblPr):
        if child.tag == qn_w("w:tblBorders"):
            tblPr.remove(child)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn_w("w:val"), "single")
        el.set(qn_w("w:sz"), sz)
        el.set(qn_w("w:space"), "0")
        el.set(qn_w("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def set_repeat_header(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    trPr.append(flag)
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def prevent_row_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, text, *, bold=False, color=INK, size=10, align="left", fill=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)
    set_cell_margins(cell)
    para = cell.paragraphs[0]
    para.clear()
    para.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.08
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_page_field(paragraph, kind: str) -> None:
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn_w("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn_w("xml:space"), "preserve")
    instr.text = f" {kind} "
    end = OxmlElement("w:fldChar")
    end.set(qn_w("w:fldCharType"), "end")
    r.append(begin)
    r.append(instr)
    r.append(end)


def set_narrow_style_fonts(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn_w("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn_w("w:ascii"), "Calibri")
    rFonts.set(qn_w("w:hAnsi"), "Calibri")
    rFonts.set(qn_w("w:eastAsia"), "Calibri")
    rFonts.set(qn_w("w:cs"), "Calibri")


def heading(doc, text, level=1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn_w("w:val"), str(level - 1))
    pPr.append(outline)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True, color=NAVY_RGB)
    else:
        set_run_font(run, size=13, bold=True, color=NAVY_RGB)


def body(doc, text, *, italic=False, muted=False, size=11, space_after=8) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, italic=italic, color=MUTED if muted else INK)


def bullet(doc, text) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=11)


def gold_rule(doc) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, GOLD)
    set_cell_text(cell, "", size=2)
    for para in cell.paragraphs:
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
    tbl = table._tbl
    tblPr = tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn_w("w:w"), "5000")
    width.set(qn_w("w:type"), "pct")
    tblPr.append(width)
    # collapse height
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn_w("w:val"), "80")
    trHeight.set(qn_w("w:hRule"), "exact")
    trPr.append(trHeight)
    set_table_borders(table, color=GOLD, sz="0")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    spacer.paragraph_format.space_before = Pt(0)


def make_table(doc, headers, rows, *, col_widths=None, row_fills=None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_header(table.rows[0])

    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = OxmlElement("w:tblW")
    tblW.set(qn_w("w:w"), "5000")
    tblW.set(qn_w("w:type"), "pct")
    tblPr.append(tblW)

    if col_widths:
        for i, cm in enumerate(col_widths):
            for cell in table.columns[i].cells:
                cell.width = Cm(cm)

    for i, title in enumerate(headers):
        set_cell_text(table.cell(0, i), title, bold=True, color=WHITE, size=10, fill=HEADER_FILL)

    for r, row in enumerate(rows, start=1):
        fill = ROW_ALT if r % 2 == 0 else "FFFFFF"
        if row_fills and r - 1 < len(row_fills) and row_fills[r - 1]:
            fill = row_fills[r - 1]
        prevent_row_split(table.rows[r])
        for c, value in enumerate(row):
            set_cell_text(table.cell(r, c), str(value), size=10, fill=fill)

    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(2)
    after.paragraph_format.space_after = Pt(8)


def scan_docs() -> dict:
    files = sorted(DOCS.rglob("*.md"))
    groups: dict[str, list[Path]] = {}
    inventory = []
    total = 0
    for path in files:
        rel = path.relative_to(DOCS)
        size = path.stat().st_size
        total += size
        group = rel.parts[0] if len(rel.parts) > 1 else "00-index"
        groups.setdefault(group, []).append(path)
        mtime = date.fromtimestamp(path.stat().st_mtime)
        inventory.append((str(rel).replace("\\", "/"), size, mtime.strftime("%d/%m/%Y")))

    group_bytes = {k: sum(p.stat().st_size for p in v) for k, v in groups.items()}

    p1 = (DOCS / "06-acceptance" / "phase-1.md").read_text(encoding="utf-8")
    acceptance_ids = re.findall(r"^\| ([A-Z]+\d+) \|", p1, flags=re.M)

    schema = (DOCS / "05-data" / "schema.md").read_text(encoding="utf-8")
    api_rows = re.findall(r"^\| (GET|POST|PUT|PATCH|DELETE) \|", schema, flags=re.M)
    tables = re.findall(r"^## \d+\. `([a-z_]+)`", schema, flags=re.M)

    return {
        "file_count": len(files),
        "total_bytes": total,
        "total_kb": round(total / 1024),
        "inventory": inventory,
        "group_bytes": group_bytes,
        "acceptance_count": len(acceptance_ids),
        "api_count": len(api_rows),
        "table_count": len(tables),
        "tables": tables,
    }


def setup_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    # First page header empty
    section.first_page_header.paragraphs[0].text = ""
    fp_footer = section.first_page_footer.paragraphs[0]
    fp_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp_footer.add_run("Nội bộ · Work OS · Không phát hành ngoài nhóm dự án")
    set_run_font(run, size=9, color=MUTED, italic=True)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = header.add_run("WORK OS")
    set_run_font(r1, size=9, bold=True, color=NAVY_RGB)
    r2 = header.add_run("  ·  Báo cáo tuần Design  ·  28/09/2026 – 04/10/2026")
    set_run_font(r2, size=9, color=MUTED)

    # gold underline via bottom border on header paragraph
    pPr = header._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn_w("w:val"), "single")
    bottom.set(qn_w("w:sz"), "12")
    bottom.set(qn_w("w:space"), "6")
    bottom.set(qn_w("w:color"), GOLD)
    pBdr.append(bottom)
    pPr.append(pBdr)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    f1 = footer.add_run("Nguồn: docs/01-product/status.md  ·  Audit 04/10/2026")
    set_run_font(f1, size=9, color=MUTED)
    footer.add_run("\t")
    f2 = footer.add_run("Trang ")
    set_run_font(f2, size=9, color=MUTED)
    add_page_field(footer, "PAGE")
    f3 = footer.add_run(" / ")
    set_run_font(f3, size=9, color=MUTED)
    add_page_field(footer, "NUMPAGES")

    # tab stop to right for page number
    tab_stops = footer.paragraph_format.tab_stops
    usable = section.page_width - section.left_margin - section.right_margin
    tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)

    pPr = footer._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn_w("w:val"), "single")
    top.set(qn_w("w:sz"), "6")
    top.set(qn_w("w:space"), "6")
    top.set(qn_w("w:color"), LINE)
    pBdr.append(top)
    pPr.append(pBdr)


def cover(doc: Document, stats: dict) -> None:
    banner = doc.add_table(rows=1, cols=1)
    set_table_borders(banner, color=NAVY, sz="0")
    cell = banner.cell(0, 0)
    shade_cell(cell, NAVY)
    set_cell_margins(cell, top=200, bottom=200, left=160, right=160)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("WORK OS")
    set_run_font(r, size=12, bold=True, color=WHITE)
    p2 = cell.add_paragraph()
    r2 = p2.add_run("BÁO CÁO TUẦN")
    set_run_font(r2, size=28, bold=True, color=WHITE)
    p3 = cell.add_paragraph()
    r3 = p3.add_run("Hạng mục Design  ·  Khóa Phase 1 và Phase 2")
    set_run_font(r3, size=12, color=RGBColor(0xD6, 0xDE, 0xE8))
    for para in cell.paragraphs:
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.space_before = Pt(0)

    gold_rule(doc)

    meta = [
        ("Sản phẩm", "Work OS — Personal + Small Team Work Operating System"),
        ("Kỳ báo cáo", "28/09/2026 – 04/10/2026  (ISO tuần 40)"),
        ("Ngày audit / xuất", "04/10/2026"),
        ("Trạng thái", "Đã khóa — không mở module design mới"),
        ("Phạm vi", "docs/  (01-product → 06-acceptance)"),
        ("Phiên bản tài liệu", "1.0"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    set_table_borders(table, color=LINE)
    for i, (k, v) in enumerate(meta):
        fill = ROW_ALT if i % 2 else "FFFFFF"
        set_cell_text(table.cell(i, 0), k, bold=True, size=10, fill=fill, color=NAVY_RGB)
        set_cell_text(table.cell(i, 1), v, size=10, fill=fill)
        table.cell(i, 0).width = Cm(4.6)
        table.cell(i, 1).width = Cm(12.0)
        prevent_row_split(table.rows[i])

    doc.add_paragraph()

    heading(doc, "1. Tóm tắt điều hành", 1)
    body(
        doc,
        "Tuần này hoàn tất audit và khóa toàn bộ design. Phase 1 (Personal daily loop) "
        "và Phase 2 (Small Team) đã đủ spec. Phase 3 (SaaS) và Phase 4 (AI) chỉ khóa hướng, "
        "không lock vendor hay chi tiết triển khai. Không còn hạng mục design đang mở.",
    )
    body(
        doc,
        f"Kết luận lấy từ docs/01-product/status.md (Audit {AUDIT.strftime('%d/%m/%Y')}): "
        "“Design Phase 1 + Phase 2 + khung Phase 3/4 đã đủ. Không mở module mới.” "
        "Cây thư mục web/ và api/ đã có skeleton Phase 1 — đó là implement, không tính vào phần design.",
    )

    heading(doc, "2. Chỉ số tuần", 1)
    kpi = [
        ("File markdown đã khóa", str(stats["file_count"])),
        ("Dung lượng spec", f"{stats['total_kb']} KB"),
        ("Nhóm tài liệu", "6  (01 → 06)"),
        ("Acceptance Phase 1", str(stats["acceptance_count"])),
        ("Bảng DB Phase 1", str(stats["table_count"])),
        ("API Phase 1", str(stats["api_count"])),
        ("Module design còn mở", "0"),
        ("Ngày ghi file docs", "04/10/2026 (toàn bộ)"),
    ]
    make_table(doc, ["Chỉ số", "Giá trị"], kpi, col_widths=[10.5, 6.1])
    body(
        doc,
        "Số file, dung lượng, số acceptance và số API được đếm trực tiếp từ docs/ "
        "lúc xuất file — không lấy từ bộ nhớ phiên chat.",
        italic=True,
        muted=True,
        size=10,
    )
    doc.add_page_break()


def content(doc: Document, stats: dict) -> None:
    heading(doc, "3. Sáu nhóm tài liệu đã khóa", 1)
    body(
        doc,
        "Thứ tự đọc khi implement: 01 → 02 → 03 → 04 → 05 → 06. "
        "Khi xung đột: domain lock thắng UX mock; UX mock thắng index. "
        "Không implement theo mock ASCII trong vision.md.",
    )
    make_table(
        doc,
        ["Nhóm", "Câu hỏi", "File đọc khi build Phase 1", "Mức khóa"],
        [
            ["01 Product", "Làm gì, khi nào?", "vision.md · status.md", "Vision — không code theo ASCII"],
            ["02 UX", "User thấy gì?", "screens-and-flows.md · visual-system.md", "Lock UI + copy English"],
            ["03 Domain", "Luật là gì?", "phase-1-rules.md · phase-1-edges.md", "Thắng khi lệch UX"],
            ["04 Architecture", "Các phần nối thế nào?", "system.md", "Lock stack + folder"],
            ["05 Data", "Lưu gì, migrate thế nào?", "schema.md · management.md", f"{stats['table_count']} bảng + {stats['api_count']} API"],
            ["06 Acceptance", "Màn nào được tính xong?", "phase-1.md", f"{stats['acceptance_count']} hàng + dogfood"],
        ],
        col_widths=[3.2, 3.8, 5.6, 4.0],
        row_fills=[ROW_INFO, ROW_OK, ROW_OK, ROW_OK, ROW_OK, ROW_OK],
    )

    heading(doc, "4. Khối lượng spec theo nhóm", 1)
    body(doc, "Dung lượng file markdown trong từng nhóm, đo lúc xuất.", muted=True, size=10)
    order = [
        ("03-domain", "03 Domain"),
        ("01-product", "01 Product"),
        ("02-ux", "02 UX"),
        ("05-data", "05 Data"),
        ("04-architecture", "04 Architecture"),
        ("06-acceptance", "06 Acceptance"),
        ("00-index", "docs/README.md"),
    ]
    rows = []
    for key, label in order:
        b = stats["group_bytes"].get(key, 0)
        if b == 0 and key != "00-index":
            continue
        share = (b / stats["total_bytes"] * 100) if stats["total_bytes"] else 0
        rows.append([label, f"{b / 1024:.1f}", f"{share:.1f}%"])
    make_table(doc, ["Nhóm", "Dung lượng (KB)", "Tỷ trọng"], rows, col_widths=[8.0, 4.3, 4.3])

    heading(doc, "5. Việc đã chốt trong tuần", 1)
    body(
        doc,
        "Tái tổ chức docs/ thành sáu nhóm. Vòng dùng hàng ngày được khóa: "
        "Inbox → Home Focus → My Tasks → Weekly Review.",
    )
    make_table(
        doc,
        ["Việc", "Kết quả", "File thắng"],
        [
            ["Khóa vòng dùng hàng ngày", "Inbox → Home Focus → My Tasks → Weekly Review", "phase-1-rules.md"],
            ["Khóa 12 bề mặt Phase 1", "Auth, Home, Inbox, Quick Add, Task, Checklist, My Tasks, Project, Calendar, Review, Search, Settings", "screens-and-flows.md"],
            ["Khóa visual + copy English", "Thang 8px, Lucide 16, token màu; mọi string UI tiếng Anh", "visual-system.md"],
            ["Khóa edge cases", "Timezone, sort, Done, Focus ≤ 3, Focus tuần tới, archive / drop / delete", "phase-1-edges.md"],
            [f"Khóa schema + API", f"{stats['table_count']} bảng Phase 1 ({', '.join(stats['tables'])}); {stats['api_count']} endpoint; không tạo bảng Phase 2", "schema.md"],
            ["Khóa vận hành DB", "Alembic, tenancy, backup 7 daily + 4 weekly, anti-pattern", "management.md"],
            ["Khóa kiến trúc", "Modular monolith: Next.js 15 + FastAPI + PostgreSQL 16", "system.md"],
            [f"Khóa acceptance", f"{stats['acceptance_count']} hàng setup → action → expected + cổng dogfood 2 tuần", "phase-1.md"],
            ["Spec Phase 2 đầy đủ", "Team, Notes, Calendar event — chỉ mở sau dogfood", "phase-2-*.md"],
            ["Khung Phase 3 / 4", "Free Personal vs Team; AI chỉ gợi ý, không tự pin", "phase-3-saas.md · phase-4-ai.md"],
        ],
        col_widths=[4.4, 8.4, 3.8],
        row_fills=[ROW_OK] * 8 + [ROW_INFO, ROW_ALT],
    )

    heading(doc, "6. Phase 1 — mười hai bề mặt đã lock", 1)
    body(doc, "Đủ để implement. Một màn khớp UX nhưng trượt một hàng acceptance thì chưa xong.")
    make_table(
        doc,
        ["Bề mặt", "Phạm vi đã khóa", "Acceptance"],
        [
            ["Auth tối thiểu", "Register / login, workspace Personal, cookie 30 ngày; không quên mật khẩu, không OAuth", "A1–A8"],
            ["Home", "Today counts + My Focus tối đa 3 + This Week Thứ Hai–Thứ Sáu", "H1–H15"],
            ["Quick Add", "Title + project + due + priority; có trên mọi màn hình app", "Q1–Q6"],
            ["Inbox", "Capture; rời inbox khi có project hoặc due; không quay lại inbox khi xóa field", "I1–I8"],
            ["Task + checklist", "Chỉ field sống; tick hết checklist không đổi status task", "T1–T13"],
            ["My Tasks", "Today / Upcoming / Overdue / All · List + Kanban bốn cột", "M1–M10"],
            ["Project", "Overview + Tasks; xóa project không xóa task", "P1–P9"],
            ["Calendar", "Due date, kéo đổi ngày; không khung giờ / meeting", "C1–C10"],
            ["Weekly Review", "Số liệu + reschedule / drop / chọn Focus tuần tới", "R1–R11"],
            ["Search · Settings", "Keyword trên title; tên, email, timezone, tên workspace", "S1–S11"],
            ["Data / API", f"{stats['table_count']} bảng, workspace-scoped, không Redis / S3", "D1–D6"],
            ["Dogfood gate", "Tự dùng ≥ 2 tuần; G1–G4 mới được sang Phase 2", "G1–G4"],
        ],
        col_widths=[3.6, 10.2, 2.8],
    )

    heading(doc, "7. Phase 2 — đã có spec, chưa build", 1)
    body(
        doc,
        "Chỉ mở khi Phase 1 đạt success: tự dùng ít nhất hai tuần; Home trả lời được "
        "hôm nay / trễ / Focus / tuần này; Review chốt được tuần sau.",
    )
    make_table(
        doc,
        ["Thứ tự", "Hạng mục", "File"],
        [
            ["1", "Invite + Owner / Admin / Member", "phase-2-team.md"],
            ["2", "Assignee (một người)", "phase-2-team.md"],
            ["3", "Comments + mentions", "phase-2-team.md"],
            ["4", "In-app notifications", "phase-2-team.md"],
            ["5", "Team dashboard + Home activity", "phase-2-team.md"],
            ["6", "Nhiều workspace + switcher", "phase-2-team.md"],
            ["7", "Calendar events (giờ, không phải due-date)", "phase-2-calendar.md"],
            ["8", "Notes (không clone Notion)", "phase-2-notes.md"],
            ["9", "Attachments + Project Files", "phase-2-team.md"],
            ["10", "Guest", "phase-2-team.md"],
            ["11", "Tags (làm cùng comments nếu rẻ)", "phase-2-team.md"],
        ],
        col_widths=[2.2, 9.6, 4.8],
    )

    heading(doc, "8. Cố ý không thiết kế", 1)
    body(doc, "Không viết spec. Không lẻn vào Phase 2.")
    make_table(
        doc,
        ["Nhóm", "Không làm"],
        [
            ["Task model", "Subtasks, dependencies, followers, multi-assignee"],
            ["Workflow", "Custom status, Gantt, project timeline"],
            ["Sản phẩm khác", "Chat, CRM, time tracking, hardware"],
            ["Kênh", "Email digest, mobile push, native app"],
            ["Doanh nghiệp", "OAuth bắt buộc, SSO, SCIM, khung i18n"],
            ["Lặp lại", "Recurring task / recurring event"],
            ["AI sớm", "Semantic search, natural-language Quick Add"],
        ],
        col_widths=[4.2, 12.4],
    )

    heading(doc, "9. Kết luận và tuần tới", 1)
    heading(doc, "9.1 Design — không viết thêm", 2)
    bullet(doc, "Không mở module mới.")
    bullet(doc, "Không bổ sung object Phase 1.")
    bullet(doc, "Không đưa subtasks, Gantt hay chat vào Phase 2.")
    heading(doc, "9.2 Implement Phase 1", 2)
    bullet(doc, "Code theo 02-ux + 03-domain + 05-data + 06-acceptance/phase-1.md.")
    bullet(doc, "Không code theo mock vision trong vision.md.")
    bullet(doc, "Phase 1 chỉ đóng khi tự dùng ≥ 2 tuần và G1–G4 đạt.")

    heading(doc, "Phụ lục A — Danh mục file docs/", 1)
    body(
        doc,
        "Mỗi dòng là một file .md trong docs/ tại thời điểm xuất. "
        "Ngày sửa lấy từ LastWriteTime trên đĩa.",
        muted=True,
        size=10,
    )
    appendix = []
    for rel, size, mtime in stats["inventory"]:
        appendix.append([rel, f"{size / 1024:.1f} KB", mtime])
    make_table(
        doc,
        ["Đường dẫn", "Dung lượng", "Ngày sửa"],
        appendix,
        col_widths=[10.6, 3.0, 3.0],
    )

    heading(doc, "Phụ lục B — Phương pháp xuất", 1)
    bullet(doc, "File này được sinh từ docs/reports/export_weekly_design_docx.py.")
    bullet(doc, "Chỉ số file / KB / acceptance / API được quét lại mỗi lần chạy.")
    bullet(doc, "Git CLI không có trên PATH trong môi trường xuất — không đối chiếu commit hash.")
    bullet(doc, "Mở file trong Microsoft Word để field số trang tự cập nhật.")


def set_core_props(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Báo cáo tuần — Design | Work OS | 28/09/2026–04/10/2026"
    props.subject = "Weekly design lock report"
    props.author = "Work OS"
    props.category = "Weekly report"
    props.comments = "Audit 2026-10-04. Phase 1 + Phase 2 design locked."
    props.language = "vi-VN"


def main() -> Path:
    stats = scan_docs()
    doc = Document()
    set_narrow_style_fonts(doc)
    setup_header_footer(doc)
    set_core_props(doc)
    cover(doc, stats)
    content(doc, stats)

    name = f"Bao-cao-tuan-Design_{WEEK_START.isoformat()}_{WEEK_END.isoformat()}.docx"
    out = OUT_DIR / name
    doc.save(out)
    print(f"Wrote {out}")
    print(
        f"files={stats['file_count']} kb={stats['total_kb']} "
        f"acceptance={stats['acceptance_count']} api={stats['api_count']} "
        f"tables={stats['table_count']}"
    )
    return out


if __name__ == "__main__":
    main()
