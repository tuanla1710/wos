import re
from datetime import date
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from sqlalchemy.orm import Session

from app.models.focus import TaskWeekTarget
from app.models.task import Task
from app.models.user import User
from app.models.workspace import Workspace
from app.services.review import _alive, review_payload
from app.services.serialize import task_out

INK = RGBColor(0x11, 0x11, 0x11)
HEADER_FILL = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Malgun Gothic"
_HANGUL = "가나다라마바사아자차카타파하"
_LEADING_LABEL = re.compile(r"^(?:\d+|[가-힣])\.\s*")


def _plain_name(name: str) -> str:
    cleaned = _LEADING_LABEL.sub("", name or "").strip()
    return cleaned or name or "No project"


def _layer_label(depth: int, index: int) -> str:
    if depth == 0:
        return f"{index}."
    n = index - 1
    if n < len(_HANGUL):
        return f"{_HANGUL[n]}."
    cycle, rest = divmod(n, len(_HANGUL))
    return f"{_HANGUL[rest]}{cycle + 1}."


def _due_text(iso: str | None) -> str:
    """Deadline as month/day, for example 10/30."""
    if not iso:
        return ""
    parts = str(iso)[:10].split("-")
    if len(parts) != 3:
        return ""
    return f"{int(parts[1]):02d}/{int(parts[2]):02d}"


def _pair(value, due: str | None) -> str:
    pct = "—" if value is None else f"{int(value)}%"
    when = _due_text(due)
    if not when:
        return pct
    return f"{pct}, {when}"


def _stored_targets(db: Session, workspace: Workspace, week_start: date) -> dict[str, int]:
    rows = (
        db.query(TaskWeekTarget)
        .filter(TaskWeekTarget.workspace_id == workspace.id, TaskWeekTarget.week_start == week_start)
        .all()
    )
    return {str(row.task_id): int(row.target_progress) for row in rows if row.target_progress is not None}


def _merge_projects(left: list[dict], right: list[dict]) -> list[dict]:
    """One project tree. A task that belongs in either week is listed on both sides."""
    index: dict[str, dict] = {}
    order: list[str] = []

    def absorb(projects: list[dict], side: str) -> None:
        for project in projects:
            pid = project.get("id") or ""
            node = index.get(pid)
            if node is None:
                node = {
                    "id": project.get("id"),
                    "name": project.get("name") or "No project",
                    "due_date": None,
                    "progress_this": None,
                    "progress_next": None,
                    "tasks": [],
                    "seen": set(),
                    "child_left": [],
                    "child_right": [],
                }
                index[pid] = node
                order.append(pid)
            if project.get("due_date"):
                node["due_date"] = project["due_date"]
            if side == "this":
                node["progress_this"] = project.get("progress")
                node["child_left"] = project.get("children") or []
            else:
                node["progress_next"] = project.get("progress")
                node["child_right"] = project.get("children") or []
            for task in project.get("tasks") or []:
                tid = task.get("id")
                if tid not in node["seen"]:
                    node["seen"].add(tid)
                    node["tasks"].append(task)

    absorb(left, "this")
    absorb(right, "next")
    merged: list[dict] = []
    for pid in order:
        node = index[pid]
        node["children"] = _merge_projects(node.pop("child_left"), node.pop("child_right"))
        merged.append(node)
    return merged


def _report_columns(
    projects: list[dict],
    this_targets: dict[str, int],
    next_targets: dict[str, int],
    limit: int = 60,
) -> tuple[list[tuple[str, bool]], list[tuple[str, bool]], list[tuple[str, bool]]]:
    """Shared task names, then this week's progress and next week's plan, each with the deadline."""
    body: list[tuple[str, bool]] = []
    this_marks: list[tuple[str, bool]] = []
    next_marks: list[tuple[str, bool]] = []

    def add(text: str, bold: bool, this_mark: str, next_mark: str) -> bool:
        if len(body) >= limit:
            return False
        body.append((text, bold))
        this_marks.append((this_mark, False))
        next_marks.append((next_mark, False))
        return True

    def walk_tasks(tasks: list[dict], depth: int) -> bool:
        for task in tasks:
            pad = "  " * depth
            if not add(f"{pad}- {task.get('title') or ''}", False, "", ""):
                return False
            if not walk_tasks(task.get("subtasks") or [], depth + 1):
                return False
        return True

    def project_percents(tasks: list[dict]) -> tuple[int | None, int | None]:
        this_vals: list[int] = []
        next_vals: list[int] = []

        def collect(items: list[dict]) -> None:
            for task in items:
                tid = str(task.get("id") or "")
                if tid in this_targets:
                    this_vals.append(this_targets[tid])
                elif task.get("progress") is not None:
                    this_vals.append(int(task["progress"]))
                if tid in next_targets:
                    next_vals.append(next_targets[tid])
                collect(task.get("subtasks") or [])

        def average(vals: list[int]) -> int | None:
            if not vals:
                return None
            return round(sum(vals) / len(vals))

        collect(tasks)
        return average(this_vals), average(next_vals)

    def walk_projects(nodes: list[dict], depth: int) -> bool:
        visible = [project for project in nodes if _has_tasks(project)]
        if depth == 0:
            def top_key(project: dict) -> tuple:
                name = project.get("name") or ""
                match = re.match(r"^(\d+)\.", name)
                if match:
                    return (0, int(match.group(1)), name)
                return (1, 0, name)

            visible.sort(key=top_key)
        else:
            visible.sort(key=lambda project: project.get("name") or "")
        shown = 0
        for project in visible:
            pad = "  " * depth
            due = project.get("due_date")
            this_pct, next_pct = project_percents(project.get("tasks") or [])
            if this_pct is None:
                this_pct = project.get("progress_this")
            if next_pct is None:
                next_pct = project.get("progress_next")
            raw_name = project.get("name") or "No project"
            if project.get("id"):
                shown += 1
                title = f"{pad}{_layer_label(depth, shown)} {_plain_name(raw_name)}"
            else:
                title = f"{pad}{raw_name}"
            if not add(
                title,
                True,
                _pair(this_pct, due),
                _pair(next_pct, due),
            ):
                return False
            if not walk_tasks(project.get("tasks") or [], depth):
                return False
            if not walk_projects(project.get("children") or [], depth + 1):
                return False
        return True

    walk_projects(projects, 0)
    if not body:
        body.append(("-", False))
        this_marks.append(("", False))
        next_marks.append(("", False))
    return body, this_marks, next_marks


def _issue_lines(overdue: list[dict], blocked: list[dict], active: list[dict]) -> tuple[list[tuple[str, bool]], list[tuple[str, bool]]]:
    issues: list[tuple[str, bool]] = []
    actions: list[tuple[str, bool]] = []
    for task in overdue[:4]:
        issues.append((f"- Overdue: {task.get('title') or ''}", False))
        actions.append((f"- Reschedule {task.get('title') or ''}", False))
    for task in blocked[:4]:
        reason = task.get("blocked_reason") or ""
        label = f"- Blocked: {task.get('title') or ''}"
        if reason:
            label += f" ({reason})"
        issues.append((label, False))
        actions.append((f"- Unblock {task.get('title') or ''}", False))
    if not issues:
        for task in active[:4]:
            issues.append((f"- In progress: {task.get('title') or ''}", False))
            actions.append(("- Continue", False))
    if not issues:
        issues.append(("-", False))
        actions.append(("-", False))
    return issues, actions


def _has_tasks(project: dict) -> bool:
    if project.get("tasks"):
        return True
    return any(_has_tasks(child) for child in project.get("children") or [])


def _set_run(run, text: str, *, size: int, bold: bool, color=INK, keep_case: bool = False) -> None:
    del bold
    run.text = text if keep_case else text.lower()
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = rPr.makeelement(qn("a:latin"), {})
        rPr.append(latin)
    latin.set("typeface", FONT)


def _fill(cell, rgb: RGBColor) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb


def _border(cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith(("lnL", "lnR", "lnT", "lnB")):
            tcPr.remove(child)
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        ln = tcPr.makeelement(qn(f"a:{edge}"), {"w": "12700", "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": "000000"})
        fill.append(clr)
        ln.append(fill)
        dash = ln.makeelement(qn("a:prstDash"), {"val": "solid"})
        ln.append(dash)
        tcPr.append(ln)


def _wrapped_count(lines: list[tuple[str, bool]], size: int, width_in: float) -> int:
    # Korean glyphs are about as wide as the font size, so the estimate stays inside the cell.
    per_line = max(6, int(width_in * 72 / size))
    total = 0
    for text, _bold in lines:
        total += max(1, (len(text) + per_line - 1) // per_line)
    return total


def _fit_font(lines: list[tuple[str, bool]], width_in: float, row_in: float, preferred: int) -> int:
    room = row_in * 72 - 4
    size = preferred
    while size > 8 and _wrapped_count(lines, size, width_in) * (size + 1) > room:
        size -= 1
    return size


def _write(cell, lines: list[tuple[str, bool]], *, size: int, align="left", anchor=MSO_ANCHOR.TOP, keep_case: bool = False) -> None:
    cell.vertical_anchor = anchor
    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    alignment = PP_ALIGN.CENTER if align == "center" else PP_ALIGN.LEFT
    for i, (text, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alignment
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 0.9
        p.level = 0
        if p.runs:
            _set_run(p.runs[0], text, size=size, bold=bold, keep_case=keep_case)
        else:
            _set_run(p.add_run(), text, size=size, bold=bold, keep_case=keep_case)


def build_review_pptx(db: Session, workspace: Workspace, user: User, week_start: date | None) -> tuple[bytes, str]:
    data = review_payload(db, workspace, user, week_start)
    today = date.fromisoformat(data["week_end"])
    active = [
        task_out(t)
        for t in _alive(db, workspace)
        .filter(Task.is_inbox.is_(False), Task.status.in_(("todo", "in_progress")))
        .order_by(Task.due_date.asc())
        .all()
        if not t.due_date or t.due_date <= today
    ]

    this = data["board"]["this"]
    nxt = data["board"]["next"]
    this_targets = _stored_targets(db, workspace, date.fromisoformat(this["start"]))
    next_targets = _stored_targets(db, workspace, date.fromisoformat(nxt["start"]))
    names, this_marks, next_marks = _report_columns(_merge_projects(this["projects"], nxt["projects"]), this_targets, next_targets)
    issues, actions = _issue_lines(data["overdue"], data["blocked"], active)

    prs = Presentation()
    slide_w = Inches(13.333333)
    slide_h = Inches(7.5)
    prs.slide_width = slide_w
    prs.slide_height = slide_h
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    margin = Inches(0.08)
    frame_w = slide_w - margin * 2
    frame_h = slide_h - margin * 2
    title_h = Inches(0.38)
    head_h = Inches(0.32)
    foot_h = Inches(0.7)
    body_h = frame_h - title_h - head_h - foot_h
    due_w = Inches(1.32)
    name_w = (frame_w - due_w * 2) // 2
    name_w2 = frame_w - due_w * 2 - name_w

    table = slide.shapes.add_table(4, 4, margin, margin, frame_w, frame_h).table
    table.columns[0].width = name_w
    table.columns[1].width = due_w
    table.columns[2].width = name_w2
    table.columns[3].width = due_w
    table.rows[0].height = title_h
    table.rows[1].height = head_h
    table.rows[2].height = body_h
    table.rows[3].height = foot_h

    table.cell(0, 0).merge(table.cell(0, 3))
    table.cell(3, 0).merge(table.cell(3, 1))
    table.cell(3, 2).merge(table.cell(3, 3))

    inset = Emu(28000)
    for row in table.rows:
        for cell in row.cells:
            cell.margin_left = inset
            cell.margin_right = inset
            cell.margin_top = Emu(14000)
            cell.margin_bottom = Emu(14000)
            _border(cell)
            _fill(cell, WHITE)

    title = table.cell(0, 0)
    _fill(title, HEADER_FILL)
    week_no = date.fromisoformat(this["start"]).isocalendar().week
    _write(title, [(f"ORAN 템팀 주간보고 W{week_no}", True)], size=19, align="center", anchor=MSO_ANCHOR.MIDDLE, keep_case=True)

    headers = (
        [("실적", True)],
        [("일시", True)],
        [("계획", True)],
        [("일시", True)],
    )
    for i, lines in enumerate(headers):
        cell = table.cell(1, i)
        _fill(cell, HEADER_FILL)
        _write(cell, lines, size=12, align="center", anchor=MSO_ANCHOR.MIDDLE)

    body_size = _fit_font(names, name_w / 914400 - 0.08, body_h / 914400 - 0.06, 12)
    mark_size = max(8, body_size - 1)
    _write(table.cell(2, 0), names, size=body_size)
    _write(table.cell(2, 1), this_marks, size=mark_size, align="center")
    _write(table.cell(2, 2), names, size=body_size)
    _write(table.cell(2, 3), next_marks, size=mark_size, align="center")

    _write(table.cell(3, 0), [("진행 업무 및 이슈", True), *issues], size=11)
    _write(table.cell(3, 2), [("대책", True), *actions], size=11)

    prs.core_properties.title = f"ORAN 템팀 주간보고 W{week_no}"
    prs.core_properties.author = user.name

    buf = BytesIO()
    prs.save(buf)
    filename = f"Weekly-Report_{data['week_start']}_{data['week_end']}.pptx"
    return buf.getvalue(), filename


def build_review_docx(db: Session, workspace: Workspace, user: User, week_start: date | None) -> tuple[bytes, str]:
    return build_review_pptx(db, workspace, user, week_start)
